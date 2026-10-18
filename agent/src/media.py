"""Conversation media cards and bounded, on-demand artifact generation."""

import asyncio
import base64
import io
import ipaddress
import json
import os
import re
import socket
import time
from contextlib import suppress
from urllib.parse import urlsplit
from uuid import uuid4

import httpx

from attachments import download_safely


def public_url(value):
    if not isinstance(value, str) or len(value) > 4096:
        return False
    try:
        url = urlsplit(value)
        host = (url.hostname or "").lower().rstrip(".")
        if (
            url.scheme != "https"
            or not host
            or url.username
            or url.password
            or url.port not in (None, 443)
        ):
            return False
        if host == "localhost" or host.endswith((".localhost", ".local", ".internal")):
            return False
        try:
            return ipaddress.ip_address(host).is_global
        except ValueError:
            return "." in host
    except ValueError:
        return False


def document_bytes(title, body):
    from docx import Document

    if not title.strip() or len(title) > 150 or not body.strip() or len(body) > 30000:
        raise ValueError("Document requires a title and at most 30,000 characters.")
    doc = Document()
    doc.add_heading(title, 0)
    for paragraph in body.split("\n"):
        doc.add_paragraph(paragraph)
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


class MediaStore:
    def __init__(self, room, attachments):
        self.room = room
        self.db = attachments.db
        self.s3 = attachments.s3

    def save(self, card, content, mime, extension, preview=""):
        key = f"generated/{self.room}/{card['id']}.{extension}"
        self.s3.put_object(
            Bucket=os.environ["ATTACHMENTS_BUCKET"],
            Key=key,
            Body=content,
            ContentType=mime,
        )
        item = {
            "room": {"S": self.room},
            "id": {"S": card["id"]},
            "key": {"S": key},
            "type": {"S": mime},
            "name": {"S": card["title"] + "." + extension},
            "size": {"N": str(len(content))},
            "expires": {"N": str(int(time.time()) + 86400)},
        }
        self.db.put_item(TableName=os.environ["ATTACHMENTS_TABLE"], Item=item)
        # URLs are fetched by the authenticated browser on click, never public.
        return {
            **card,
            "status": "ready",
            "assetId": card["id"],
            "preview": preview.encode("utf-8")[:6000].decode("utf-8", errors="ignore"),
        }

    def uploaded(self, attachment_id):
        from attachments import attachment_key

        key = attachment_key(self.room, attachment_id)
        item = self.db.get_item(
            TableName=os.environ["ATTACHMENTS_TABLE"],
            Key={"room": {"S": self.room}, "id": {"S": attachment_id}},
            ConsistentRead=True,
        ).get("Item")
        if (
            not item
            or int(item["expires"]["N"]) <= time.time()
            or item["key"]["S"] != key
        ):
            raise ValueError("Attachment unavailable in this conversation.")
        self.s3.head_object(Bucket=os.environ["ATTACHMENTS_BUCKET"], Key=key)
        mime = item["type"]["S"]
        return {
            "id": attachment_id,
            "assetId": attachment_id,
            "title": item["name"]["S"][:150],
            "kind": "image"
            if mime.startswith("image/")
            else "video"
            if mime.startswith("video/")
            else "document",
            "origin": "uploaded",
            "status": "ready",
        }


class GenerationClient:
    def __init__(self, key, *, transport=None):
        self.key = key
        self.transport = transport

    async def json(self, client, method, path, body=None, limit=32 * 1024 * 1024):
        async with client.stream(
            method,
            "https://openrouter.ai/api/v1/" + path,
            headers={"Authorization": f"Bearer {self.key}"},
            json=body,
        ) as response:
            response.raise_for_status()
            result = bytearray()
            async for chunk in response.aiter_bytes():
                result.extend(chunk)
                if len(result) > limit:
                    raise ValueError("Provider response too large.")
            return json.loads(result)

    async def image(self, prompt, model):
        from PIL import Image

        async with httpx.AsyncClient(transport=self.transport, timeout=120) as client:
            value = await self.json(
                client,
                "POST",
                "images",
                {"model": model, "prompt": prompt, "n": 1, "output_format": "png"},
            )
        raw = base64.b64decode(value["data"][0]["b64_json"], validate=True)
        if len(raw) > 20 * 1024 * 1024:
            raise ValueError("Generated image too large.")
        with Image.open(io.BytesIO(raw)) as image:
            if image.width * image.height > 20_000_000:
                raise ValueError("Generated image resolution too large.")
            output = io.BytesIO()
            image.convert("RGB").save(output, format="PNG")
        return output.getvalue()

    async def video(self, prompt, model):
        async with httpx.AsyncClient(
            transport=self.transport, timeout=60, follow_redirects=False
        ) as client:
            models = await self.json(
                client, "GET", "videos/models", limit=2 * 1024 * 1024
            )
            info = next((m for m in models["data"] if m["id"] == model), None)
            if (
                not info
                or 4 not in info.get("supported_durations", [])
                or "720p" not in info.get("supported_resolutions", [])
            ):
                raise ValueError(
                    "Configured video model does not support four seconds at 720p."
                )
            job = await self.json(
                client,
                "POST",
                "videos",
                {
                    "model": model,
                    "prompt": prompt,
                    "duration": 4,
                    "resolution": "720p",
                    "aspect_ratio": "16:9",
                    "generate_audio": False,
                },
            )
            job_id = job.get("id", "")
            if not re.fullmatch(r"[A-Za-z0-9_-]{1,200}", job_id):
                raise ValueError("Invalid video job identifier.")
            for _ in range(90):
                state = await self.json(
                    client, "GET", f"videos/{job_id}", limit=1024 * 1024
                )
                if state.get("status") == "completed":
                    break
                if state.get("status") in ("failed", "cancelled", "expired"):
                    raise ValueError("Video generation failed.")
                await asyncio.sleep(5)
            else:
                raise TimeoutError("Video generation took too long.")
            return await self.download_video(client, job_id)

    async def download_video(self, client, job_id):
        url = f"https://openrouter.ai/api/v1/videos/{job_id}/content?index=0"
        headers = {"Authorization": f"Bearer {self.key}"}
        for _ in range(4):
            async with client.stream("GET", url, headers=headers) as response:
                if response.is_redirect:
                    url = str(response.url.join(response.headers["location"]))
                    if not public_url(url):
                        raise ValueError("Unsafe video download URL.")
                    # Public provider redirects only. Never forward the API key to a CDN.
                    addresses = await asyncio.to_thread(
                        socket.getaddrinfo,
                        urlsplit(url).hostname,
                        443,
                        type=socket.SOCK_STREAM,
                    )
                    if not addresses or any(
                        not ipaddress.ip_address(x[4][0]).is_global for x in addresses
                    ):
                        raise ValueError("Non-public video download host.")
                    headers = {}
                    continue
                response.raise_for_status()
                data = bytearray()
                async for chunk in response.aiter_bytes():
                    data.extend(chunk)
                    if len(data) > 60 * 1024 * 1024:
                        raise ValueError("Generated video too large.")
                if len(data) < 12 or data[4:8] != b"ftyp":
                    raise ValueError("Provider did not return an MP4 video.")
                return bytes(data)
        raise ValueError("Too many video redirects.")


class Media:
    def __init__(
        self,
        room,
        publish,
        *,
        store,
        api_key="",
        image_model="",
        video_model="",
        client=None,
        announce=None,
        search=None,
    ):
        self.room = room
        self.publish = publish
        self.store = store
        self.image_model = image_model
        self.video_model = video_model
        self.client = client or GenerationClient(api_key)
        self.api_key = api_key
        self.announce = announce
        self.search = search
        self.tasks = set()
        self.count = 0
        self.closed = False
        self.pending = False
        self.playing = False

    async def emit(self, card):
        await self.publish({**card, "room": self.room})

    async def create(self, kind, title, prompt):
        if (
            kind not in ("image", "video", "document")
            or not title.strip()
            or len(title) > 150
            or not prompt.strip()
            or len(prompt) > (30000 if kind == "document" else 2000)
        ):
            return {"error": "Invalid media request."}
        if kind != "document" and (
            not self.api_key
            or not (self.image_model if kind == "image" else self.video_model)
        ):
            return {
                "error": "This generation model is not configured. Do not claim the media was created."
            }
        if self.closed or self.pending or self.tasks or self.count >= 5:
            return {
                "error": "One generation can run at a time, with up to five per conversation."
            }
        card = {
            "id": str(uuid4()),
            "kind": kind,
            "title": title,
            "origin": "generated",
            "status": "processing",
        }
        # Publish before spending credits. If delivery fails, do not start a job.
        self.pending = True
        try:
            await self.emit(card)
        except Exception:
            return {
                "error": "Could not deliver the media card. No generation was started."
            }
        finally:
            self.pending = False
        self.count += 1
        task = asyncio.create_task(self._generate(card, prompt))
        self.tasks.add(task)
        task.add_done_callback(self.tasks.discard)
        return {
            "status": "processing",
            "id": card["id"],
            "message": "The result will appear as a media card. Continue the conversation; do not claim it is ready yet.",
        }

    async def _generate(self, card, prompt):
        try:
            kind = card["kind"]
            if kind == "document":
                data = await download_safely(document_bytes, card["title"], prompt)
                mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                extension = "docx"
            elif kind == "image":
                data = await asyncio.wait_for(
                    self.client.image(prompt, self.image_model), 150
                )
                mime = "image/png"
                extension = "png"
            else:
                data = await asyncio.wait_for(
                    self.client.video(prompt, self.video_model), 600
                )
                mime = "video/mp4"
                extension = "mp4"
            card = await download_safely(
                self.store.save,
                card,
                data,
                mime,
                extension,
                prompt if kind == "document" else "",
            )
            await self.emit(card)
            if self.announce and not self.closed and not self.playing:
                self.announce(
                    f"The {kind} is ready. You can open it in the conversation."
                )
        except asyncio.CancelledError:
            raise
        except Exception:
            with suppress(Exception):
                await self.emit(
                    {
                        **card,
                        "status": "error",
                        "message": "Could not finish this media. You can ask Diana to try again.",
                    }
                )

    async def show_attachment(self, attachment_id):
        try:
            card = await download_safely(self.store.uploaded, attachment_id)
            await self.emit(card)
            return {"status": "shown", "title": card["title"]}
        except Exception:
            return {"error": "This attachment is unavailable in this conversation."}

    async def find(self, kind, query, current_time):
        if (
            kind not in ("image", "video", "document")
            or not query.strip()
            or len(query) > 300
        ):
            return {"error": "Invalid media search."}
        try:
            cards = []
            if kind == "image":
                async with httpx.AsyncClient(
                    timeout=20,
                    headers={"User-Agent": "Diana/1.0 (conversation media search)"},
                ) as client:
                    response = await client.get(
                        "https://commons.wikimedia.org/w/api.php",
                        params={
                            "action": "query",
                            "format": "json",
                            "generator": "search",
                            "gsrsearch": query,
                            "gsrnamespace": 6,
                            "gsrlimit": 3,
                            "prop": "imageinfo",
                            "iiprop": "url|extmetadata",
                            "iiurlwidth": 1000,
                        },
                    )
                    response.raise_for_status()
                    if len(response.content) > 200000:
                        raise ValueError("Response too large")
                for page in response.json().get("query", {}).get("pages", {}).values():
                    info = page.get("imageinfo", [{}])[0]
                    url = info.get("thumburl", info.get("url", ""))
                    source = info.get("descriptionurl", "")
                    if (
                        public_url(url)
                        and urlsplit(url).hostname == "upload.wikimedia.org"
                        and public_url(source)
                    ):
                        cards.append(
                            {
                                "url": url,
                                "sourceUrl": source,
                                "title": page.get("title", "Image")[:150],
                                "attribution": "Wikimedia Commons — creator and license on source page",
                            }
                        )
            elif self.search:
                suffix = (
                    " video site:youtube.com/watch"
                    if kind == "video"
                    else " filetype:pdf"
                )
                result = await self.search.search(query + suffix, current_time)
                for source in result.get("sources", [])[:3]:
                    if public_url(source["url"]):
                        cards.append(
                            {
                                "url": source["url"],
                                "sourceUrl": source["url"],
                                "title": source["title"][:150],
                            }
                        )
            for card in cards[:3]:
                await self.emit(
                    {
                        **card,
                        "id": str(uuid4()),
                        "kind": kind,
                        "origin": "web",
                        "status": "ready",
                    }
                )
            return {
                "status": "shown" if cards else "unavailable",
                "sources": cards[:3],
                "message": "These are source links, not proof their contents have been inspected. Do not invent what a video or document says.",
            }
        except Exception:
            return {"error": "Media search is unavailable. Do not invent media URLs."}

    async def close(self):
        self.closed = True
        for task in list(self.tasks):
            task.cancel()
        await asyncio.gather(*self.tasks, return_exceptions=True)
