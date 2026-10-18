import pytest

from media import Media, document_bytes, public_url


def test_public_links_reject_local_credentials_and_scripts():
    assert public_url("https://example.org/a.pdf")
    for url in [
        "http://example.org/a",
        "https://127.0.0.1/x",
        "https://169.254.169.254/",
        "https://user:pass@example.org/x",
        "javascript:alert(1)",
        "https://localhost/x",
    ]:
        assert not public_url(url)


def test_document_is_real_docx_and_bounded():
    import io

    from docx import Document

    data = document_bytes("Trip notes", "Bring a jacket.\nCall the hotel.")
    doc = Document(io.BytesIO(data))
    assert "Trip notes" in doc.paragraphs[0].text
    assert "Bring a jacket." in doc.paragraphs[1].text
    with pytest.raises(ValueError):
        document_bytes("x", "x" * 30001)


async def test_generation_unconfigured_does_not_claim_success():
    async def publish(card):
        raise AssertionError("must not publish")

    media = Media(
        "diana_room",
        publish,
        store=object(),
        api_key="",
        image_model="",
        video_model="",
    )
    result = await media.create("image", "Sunset", "A sunset")
    assert "error" in result


async def test_background_job_delivers_real_document_and_limits_parallel_work():
    import asyncio

    cards = []

    class Store:
        def save(self, card, data, mime, extension, preview):
            assert data[:2] == b"PK" and extension == "docx"
            return {
                **card,
                "status": "ready",
                "assetId": card["id"],
                "preview": preview,
            }

    async def publish(card):
        cards.append(card)

    media = Media("diana_room", publish, store=Store())
    result = await media.create("document", "Notes", "A few useful notes.")
    assert result["status"] == "processing"
    assert "error" in await media.create("document", "Duplicate", "text")
    await asyncio.gather(*media.tasks)
    assert [c["status"] for c in cards] == ["processing", "ready"]
    assert cards[-1]["room"] == "diana_room"
    assert cards[-1]["preview"] == "A few useful notes."
    await media.close()


async def test_provider_image_request_and_video_polling_do_not_trust_urls():
    import base64
    import io
    import json

    import httpx
    from PIL import Image

    from media import GenerationClient

    data = io.BytesIO()
    Image.new("RGB", (32, 32), "blue").save(data, format="PNG")
    requests = []

    def handle(request):
        requests.append(request)
        assert request.headers["authorization"] == "Bearer private-test-key"
        path = request.url.path
        if path.endswith("/images"):
            assert json.loads(request.content)["model"] == "image-model"
            return httpx.Response(
                200,
                json={
                    "data": [{"b64_json": base64.b64encode(data.getvalue()).decode()}]
                },
            )
        if path.endswith("/models"):
            return httpx.Response(
                200,
                json={
                    "data": [
                        {
                            "id": "video-model",
                            "supported_durations": [4],
                            "supported_resolutions": ["720p"],
                        }
                    ]
                },
            )
        if path.endswith("/videos"):
            body = json.loads(request.content)
            assert body["duration"] == 4 and body["generate_audio"] is False
            return httpx.Response(
                202, json={"id": "job-1", "polling_url": "https://evil.test/steal-key"}
            )
        if path.endswith("/content"):
            return httpx.Response(200, content=b"\x00\x00\x00\x18ftypmp42test")
        return httpx.Response(
            200,
            json={
                "status": "completed",
                "unsigned_urls": ["https://evil.test/steal-key"],
            },
        )

    client = GenerationClient("private-test-key", transport=httpx.MockTransport(handle))
    assert (await client.image("blue", "image-model")).startswith(b"\x89PNG")
    assert b"ftyp" in await client.video("blue", "video-model")
    assert all(request.url.host == "openrouter.ai" for request in requests)


async def test_external_video_redirect_never_receives_api_key(monkeypatch):
    import socket

    import httpx

    from media import GenerationClient

    monkeypatch.setattr(
        socket, "getaddrinfo", lambda *a, **kw: [(2, 1, 6, "", ("93.184.216.34", 443))]
    )

    def handle(request):
        if request.url.host == "openrouter.ai":
            return httpx.Response(
                302, headers={"location": "https://cdn.example.org/video.mp4"}
            )
        assert "authorization" not in request.headers
        return httpx.Response(200, content=b"\x00\x00\x00\x18ftypmp42test")

    gen = GenerationClient("private-test-key")
    async with httpx.AsyncClient(transport=httpx.MockTransport(handle)) as client:
        assert b"ftyp" in await gen.download_video(client, "job")


async def test_failure_publishes_error_not_fake_ready():
    import asyncio

    cards = []

    class Generator:
        async def image(self, *args):
            raise RuntimeError("secret-provider-error")

    async def publish(card):
        cards.append(card)

    media = Media(
        "diana_room",
        publish,
        store=object(),
        api_key="key",
        image_model="model",
        client=Generator(),
    )
    await media.create("image", "Photo", "A flower")
    await asyncio.gather(*media.tasks)
    assert [c["status"] for c in cards] == ["processing", "error"]
    assert "secret-provider-error" not in str(cards)
